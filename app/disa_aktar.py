from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .etiketler import DURUM_ETIKET

KOLONLAR = [
    ("linkedin_ad", "Kişi", 24),
    ("linkedin_rol", "Rol", 24),
    ("linkedin_unvan", "LinkedIn unvanı", 40),
    ("isletme_adi", "Şirket", 28),
    ("sirket_boyutu", "Çalışan sayısı", 15),
    ("sirket_sektoru", "Şirket sektörü", 24),
    ("konum", "Konum", 22),
    ("kampanya", "Kampanya", 24),
    ("puan", "Puan", 8),
    ("puan_detay", "Puanın gerekçesi", 50),
    ("son_paylasim_gun", "Son paylaşım (gün önce)", 14),
    ("website", "Web sitesi", 30),
    ("google_arama_terimi", "Google araması", 24),
    ("google_sirasi", "Google sırası", 13),
    ("google_sayfasi", "Google sayfası", 13),
    ("linkedin_url", "LinkedIn", 40),
    ("sirket_linkedin", "Şirketin LinkedIn sayfası", 40),
    ("durum", "Durum", 21),
    ("notlar", "Notlar", 45),
    ("son_islem_tarihi", "Son işlem", 17),
]
LINK_KOLONLARI = {"website", "linkedin_url", "sirket_linkedin"}


def _hucre_degeri(lead: dict, anahtar: str, kampanya_adlari: dict):
    deger = lead.get(anahtar)
    if anahtar == "kampanya":
        return kampanya_adlari.get(lead.get("kampanya_id"))
    if anahtar == "durum":
        return DURUM_ETIKET.get(deger, deger)
    if anahtar == "son_paylasim_gun" and deger is not None and deger < 0:
        return "Paylaşım yok"
    if anahtar == "google_sirasi" and deger is None and lead.get("google_arama_terimi"):
        return "İlk 30'da yok"
    if anahtar == "son_islem_tarihi" and deger:
        try:
            return datetime.fromisoformat(deger)
        except ValueError:
            return deger
    return deger


def excele_aktar(yol: str, leads: list, kampanya_adlari: dict = None) -> None:
    kampanya_adlari = kampanya_adlari or {}
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Adaylar"
    sayfa.append([baslik for _, baslik, _ in KOLONLAR])

    for lead in leads:
        sayfa.append([_hucre_degeri(lead, anahtar, kampanya_adlari) for anahtar, _, _ in KOLONLAR])
        satir = sayfa.max_row
        for sutun, (anahtar, _, _) in enumerate(KOLONLAR, start=1):
            hucre = sayfa.cell(row=satir, column=sutun)
            if anahtar in LINK_KOLONLARI and str(hucre.value or "").lower().startswith("http"):
                hucre.hyperlink = hucre.value
                hucre.style = "Hyperlink"
            elif anahtar == "son_islem_tarihi" and isinstance(hucre.value, datetime):
                hucre.number_format = "DD.MM.YYYY HH:MM"

    baslik_dolgu = PatternFill("solid", fgColor="4F46E5")
    for sutun, (_, _, genislik) in enumerate(KOLONLAR, start=1):
        hucre = sayfa.cell(row=1, column=sutun)
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = baslik_dolgu
        hucre.alignment = Alignment(vertical="center")
        sayfa.column_dimensions[get_column_letter(sutun)].width = genislik
    sayfa.row_dimensions[1].height = 22
    sayfa.freeze_panes = "A2"
    sayfa.auto_filter.ref = sayfa.dimensions
    kitap.save(yol)
